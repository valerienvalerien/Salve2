#!/usr/bin/env python3
"""Génère PROJECTION-FINANCES-SALVERYS.xlsx depuis tools/finance-model.mjs.

Le classeur n'est PAS un instantané : chaque cellule calculée est une vraie formule
Excel qui remonte à l'onglet « Hypothèses ». Changer le taux de change ou le nombre de
positions recalcule tout, comme dans la page HTML.

Les valeurs par défaut sont lues depuis le moteur JavaScript, pas recopiées à la main :
une modification de tools/finance-model.mjs se retrouve ici au prochain export.

    python3 tools/export-xlsx.py

Vérification : tools/verify-xlsx.py compare les cellules calculées au moteur.
"""
import json
import math
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SORTIE = ROOT / "PROJECTION-FINANCES-SALVERYS.xlsx"
LOGO = ROOT / "assets" / "logo-salverys.png"

# Charte Salverys (charte-graphique-salverys.html)
ENCRE = "2E3B4E"
VERT = "6F8770"
OR = "B89968"
SABLE = "F6F2EA"
TRAIT = "E7DFD1"
ROUGE = "B6604D"

BLEU_SAISIE = "0000FF"   # convention modèle financier : entrée modifiable
NOIR = "000000"

EUR = '#,##0 "€";(#,##0 "€");-'
EUR2 = '#,##0.00 "€";(#,##0.00 "€");-'
PCT = '0.0%'
NB = '#,##0;(#,##0);-'
AR = '#,##0 "Ar";;-'


def modele():
    """Valeurs par défaut du moteur, pour ne pas les recopier à la main."""
    js = """
import('./tools/finance-model.mjs').then(async M => {
  const { SCENARIOS } = await import('./tools/finance-scenarios.mjs');
  const deals = SCENARIOS.reference;
  process.stdout.write(JSON.stringify({
    a: M.ASSUMPTIONS, offers: M.OFFERS, entree: M.MB_ENTRY_POSITIONS,
    direct: M.DIRECT_REFERENCE, plancher: M.PRICE_FLOOR,
    agent: M.monthlyAgentCost(), manager: M.monthlyManagerCost(),
    scenarios: SCENARIOS,
    rows: M.projectCash(deals, 24),
    contrib: M.contributions(deals),
    ca: deals.map(d => M.mbRevenue(d.metier, d.positions)),
    reserveRate: M.revenueReserveRate(),
    onb: deals.map(d => M.onboardingCost(d.positions)),
  }));
});
"""
    out = subprocess.run(["node", "-e", js], cwd=ROOT, capture_output=True, text=True, check=True)
    return json.loads(out.stdout)


# ----------------------------------------------------------------- mise en forme

def logo_allege():
    """Le logo du dépôt fait 3 060 px et 100 Ko ; embarqué six fois il alourdit le
    classeur pour rien. On le réduit à la volée, sans toucher au fichier d'origine."""
    from io import BytesIO
    try:
        from PIL import Image
    except ImportError:
        return str(LOGO)
    im = Image.open(LOGO)
    im.thumbnail((640, 640), Image.LANCZOS)
    tampon = BytesIO()
    im.save(tampon, format="PNG", optimize=True)
    tampon.seek(0)
    return tampon


def logo(ws, ancre="A1"):
    """Logo Salverys en tête de chaque onglet, à hauteur constante."""
    img = XLImage(logo_allege())
    img.height = 40
    img.width = int(40 * 3060 / 780)
    ws.add_image(img, ancre)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 8


def titre(ws, texte, sous_titre=""):
    ws["C1"] = texte
    ws["C1"].font = Font(name="Arial", size=14, bold=True, color=ENCRE)
    if sous_titre:
        ws["C2"] = sous_titre
        ws["C2"].font = Font(name="Arial", size=9, italic=True, color="777777")
        ws.row_dimensions[2].height = 14


def entete(ws, ligne, libelles, largeurs=None):
    for i, lab in enumerate(libelles, start=1):
        c = ws.cell(row=ligne, column=i, value=lab)
        c.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ENCRE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side("thin", color=ENCRE))
    ws.row_dimensions[ligne].height = 28
    if largeurs:
        for i, w in enumerate(largeurs, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def bloc(ws, ligne, texte, couleur=OR):
    c = ws.cell(row=ligne, column=1, value=texte)
    c.font = Font(name="Arial", size=9, bold=True, color=ENCRE)
    c.fill = PatternFill("solid", fgColor=SABLE)
    for col in range(2, 6):
        ws.cell(row=ligne, column=col).fill = PatternFill("solid", fgColor=SABLE)
    ws.cell(row=ligne, column=1).border = Border(left=Side("thick", color=couleur))


def saisie(cell, fmt=None):
    cell.font = Font(name="Arial", size=10, color=BLEU_SAISIE, bold=True)
    cell.fill = PatternFill("solid", fgColor="FFFDE7")
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal="right")


def calcul(cell, fmt=None, gras=False, valeur=None):
    cell.font = Font(name="Arial", size=10, color=NOIR, bold=gras)
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal="right")
    if valeur is not None:
        VALEURS.setdefault(cell.parent.title, {})[cell.coordinate] = valeur


# LibreOffice Calc n'est pas installé dans cet environnement (seul libreoffice-core l'est),
# donc aucun recalcul externe n'est possible. openpyxl, lui, écrit les formules SANS valeur
# en cache : un lecteur qui ne recalcule pas — pandas, un aperçu, Google Sheets à l'import —
# verrait des cellules vides. On injecte donc la valeur calculée PAR LE MOTEUR JavaScript
# à côté de chaque formule. Les formules restent vivantes : le classeur porte
# fullCalcOnLoad, donc Excel recalcule tout à l'ouverture et écrase ces caches.
VALEURS = {}


def injecter_valeurs(chemin, wb):
    """Ajoute <v> après chaque <f> dont on connaît la valeur, dans le XML du classeur."""
    feuille_par_nom = {ws.title: f"xl/worksheets/sheet{i + 1}.xml" for i, ws in enumerate(wb.worksheets)}
    tampon = chemin.with_suffix(".tmp.xlsx")
    poses = 0
    with zipfile.ZipFile(chemin) as src, zipfile.ZipFile(tampon, "w", zipfile.ZIP_DEFLATED) as dst:
        cibles = {feuille_par_nom[n]: v for n, v in VALEURS.items() if n in feuille_par_nom}
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename in cibles:
                xml = data.decode("utf-8")
                valeurs = cibles[item.filename]

                def remplace(mo):
                    nonlocal poses
                    coord, attrs, formule = mo.group(1), mo.group(2), mo.group(3)
                    if coord not in valeurs:
                        return mo.group(0)
                    v = valeurs[coord]
                    poses += 1
                    if isinstance(v, str):
                        esc = v.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                        return f'<c r="{coord}"{attrs} t="str">{formule}<v>{esc}</v></c>'
                    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
                        return mo.group(0)
                    return f'<c r="{coord}"{attrs}>{formule}<v>{v!r}</v></c>'

                # openpyxl écrit « <f>FORMULE</f><v /> » : le <v> existe mais est vide.
                xml = re.sub(
                    r'<c r="([A-Z]+\d+)"((?:(?!/>)[^>])*)>(<f>.*?</f>)(?:<v\s*/>|<v>.*?</v>)?</c>',
                    remplace, xml, flags=re.S)
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    shutil.move(tampon, chemin)
    return poses


def texte(cell, taille=10, gras=False, couleur=NOIR, italique=False, wrap=False):
    cell.font = Font(name="Arial", size=taille, bold=gras, color=couleur, italic=italique)
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def note(ws, ligne, contenu, colonnes=6):
    c = ws.cell(row=ligne, column=1, value=contenu)
    texte(c, taille=9, couleur="777777", italique=True, wrap=True)
    ws.merge_cells(start_row=ligne, start_column=1, end_row=ligne, end_column=colonnes)
    ws.row_dimensions[ligne].height = 30


# --------------------------------------------------------------------- onglets

def onglet_hypotheses(wb, m):
    ws = wb.create_sheet("Hypothèses")
    logo(ws)
    titre(ws, "Hypothèses", "Cellules bleues : à modifier. Tout le classeur en dépend. Noir : calculé.")
    entete(ws, 4, ["Paramètre", "Valeur", "Unité", "Statut et source"], [42, 14, 14, 74])

    a, L = m["a"], {}

    def ligne(r, lab, val, unite, source, fmt=None, formule=None, valeur=None):
        texte(ws.cell(row=r, column=1, value=lab), taille=10)
        c = ws.cell(row=r, column=2, value=formule if formule else val)
        if formule:
            calcul(c, fmt, valeur=valeur)
        else:
            saisie(c, fmt)
        texte(ws.cell(row=r, column=3, value=unite), taille=9, couleur="777777")
        texte(ws.cell(row=r, column=4, value=source), taille=9, couleur="777777", wrap=True)
        L[lab] = r

    bloc(ws, 5, "PERSONNEL")
    ligne(6, "Brut mensuel manager (et agent si B39 = 0)", a["grossSalaryAr"], "Ar", "Décision direction 2026-08-14. À confirmer sur bulletins de paie. Depuis le 2026-09-24, l'agent se chiffre au coût complet de la ligne 39.", AR)
    ligne(7, "Taux de change", a["exchangeArPerEuro"], "Ar / €", "Convention de simulation. Remplacer par le taux de paiement réel et ses frais.", '#,##0')
    ligne(8, "Cotisations employeur", a["employerRate"], "% du brut", "HYPOTHÈSE RÉGLEMENTAIRE NON VÉRIFIÉE. À confirmer auprès du gestionnaire de paie.", PCT)
    ligne(9, "SME retenu", a["smeAr"], "Ar", "lexpress.mg, février 2026. Sert au plafond de cotisation.", AR)
    ligne(10, "Plafond de cotisation (8 × SME)", None, "Ar", "Calculé.", AR, "=B9*8", valeur=a["smeAr"] * 8)
    ligne(11, "VoIP par agent", a["voipPerAgent"], "€ / mois", "Devis ou facture à obtenir.", EUR)
    ligne(12, "VoIP par manager", a["voipPerManager"], "€ / mois", "Hypothèse : le manager n'a pas de ligne facturée. À corriger s'il décroche.", EUR)
    ligne(13, "Coussin sur personnel", a["contingency"], "%", "Réserve prudente, retirée du cash. Ce n'est pas une paie versée.", PCT)
    ligne(14, "Coût d'un agent", None, "€ / mois", "Calculé : coût complet (ligne 39) ÷ change + VoIP, sans coussin. Si la ligne 39 vaut 0 : (brut + cotisations plafonnées) ÷ change, plus VoIP, plus coussin.", EUR2,
          "=IF(B39>0,B39/B7+B11,((B6+MIN(B6,B10)*B8)/B7+B11)*(1+B13))", valeur=m["agent"])
    ligne(15, "Coût d'un manager", None, "€ / mois", "Calculé, même formule sans la VoIP agent.", EUR2,
          "=((B6+MIN(B6,B10)*B8)/B7+B12)*(1+B13)", valeur=m["manager"])
    ligne(16, "Positions par manager", a["managerCapacity"], "positions", "Un manager de plus par tranche, DANS CHAQUE MÉTIER.", NB)
    ligne(17, "Positions encadrées par le fondateur", a["founderSupervisesUpTo"], "positions", "Arbitrage A10 du 2026-09-24 : le fondateur encadre le lancement à 5 positions. Un manager s'ajoute à la 6e ; le forfait d'encadrement de la ligne 40 est alors à recalculer.", NB)

    bloc(ws, 18, "STRUCTURE ET DÉMARRAGE")
    ligne(19, "Outils récurrents connus", a["toolsPerMonth"], "€ / mois", "REGISTRE-COUTS-OUTILS.md : 65 € d'abonnements + 36 €/an de domaines. Montants DÉCLARÉS, non rapprochés de factures.", EUR)
    ligne(20, "Charges mensuelles supplémentaires", a["otherMonthlyCost"], "€ / mois", "À REMPLIR : taxes, comptabilité, assurance, frais bancaires, équipement, congés. Laisser 0 est une décision.", EUR)
    ligne(21, "Onboarding — socle par contrat", a["onboardingBase"], "€", "AUDIT-ESPACE-CLIENT.md §9.2. Ordre de grandeur non documenté sur temps passé.", EUR)
    ligne(22, "Onboarding — par position", a["onboardingPerPosition"], "€", "Idem. Le forfait de 1 600 €/contrat utilisé jusqu'au 2026-09-16 ne valait que pour 3 positions.", EUR)
    ligne(23, "Embauche après signature", a["hireDelayMonths"], "mois", "C'est ce délai qui crée le creux : la paie part avant la première facture.", NB)
    ligne(24, "Première facture après signature", a["firstInvoiceDelayMonths"], "mois", "Hypothèse. Le devis date le début de facturation, et c'est lui qui fait foi.", NB)

    bloc(ws, 25, "ENCAISSEMENT ET RISQUE")
    ligne(26, "Règlement après facture", a["paymentDelayMonths"], "mois", "PRICING.md §3 : paiement à 30 jours. Paramètre le plus sensible du modèle.", NB)
    ligne(27, "Dépôt par position", a["depositPerPosition"], "€", "SUPPRIMÉ de la grille le 2026-09-24 (arbitrage A1), remplacé par les frais de mise en service. Gardé pour simuler l'ancien dépôt rendu par crédit de facture.", EUR)
    ligne(28, "Crédit de dépôt par facture", a["invoiceCreditPerPosition"], "€", "Ancien mécanisme, à 0.", EUR)
    ligne(29, "Factures créditées", a["invoiceCreditMonths"], "nb", "Ancien mécanisme.", NB)
    ligne(30, "Résidu de dépôt NON rendu", None, "€ / position", "Doit valoir 0. Positif = frais d'activation déguisé, à nommer. Négatif = argent jamais encaissé, offert.", EUR, "=B27-B28*B29", valeur=a["depositPerPosition"] - a["invoiceCreditPerPosition"] * a["invoiceCreditMonths"])
    ligne(31, "Frais de mise en service par contrat", a["activationFeePerDeal"], "€", "Part fixe éventuelle. PRICING.md §3 n'en prévoit pas.", EUR)
    ligne(32, "Frais de mise en service par position", a["activationFeePerPosition"], "€", "PRICING.md §3, arbitrage A1 du 2026-09-24 : 490 €/position à la signature, acquis, non imputés sur les mensualités.", EUR)
    ligne(33, "Impayés sur CA encaissé", a["badDebtRate"], "%", "0 % n'est pas une mesure, c'est une absence d'historique.", PCT)
    ligne(34, "Banque, change et taxes sur CA encaissé", a["revenueChargeRate"], "%", "Arbitrage A9 : 1 % de frais bancaires et risque de change. Taxes non comprises : taux non vérifiés.", PCT)

    bloc(ws, 35, "DÉPART")
    ligne(36, "Trésorerie de départ", a["initialCash"], "€", "0 mesure un BESOIN THÉORIQUE. Ce n'est pas le solde réel de Salverys.", EUR)

    bloc(ws, 38, "ARBITRAGES DU 2026-09-24 (BMC v1.3)")
    ligne(39, "Coût complet agent, tout compris", a["agentAllInAr"], "Ar / mois", "Arbitrage A7 : rémunération, charges, congés, absences, formation et remplacement. À VALIDER PAR LA PAIE : le brut de la ligne 6 chargé n'y tient pas. 0 = calcul par le brut.", AR)
    ligne(40, "Structure dès la 1re embauche", a["structureMonthly"], "€ / mois", "Arbitrage A8, hypothèse pessimiste : encadrement / QA 600 + loyer 60 m² 720 + double fibre 300 + RH / paie / comptable 300 + énergie 300. Payée tant qu'un agent est en poste.", EUR)
    ligne(41, "Poste de travail par position", a["equipmentPerPosition"], "€", "Laptop + écran, 4 M Ar, décaissés à la signature.", EUR)
    ligne(42, "Groupe électrogène + UPS", a["equipmentBase"], "€", "2,25 M Ar, décaissés une fois, au premier mois avec du personnel. À tester avant d'être promis.", EUR)
    ligne(43, "Provision avoirs SLA sur CA encaissé", a["slaReserveRate"], "%", "Arbitrage A9 : avoir de 20 % au plus un mois sur douze = 1,7 % du CA ; provision de 2 %.", PCT)

    ws.cell(row=30, column=2).font = Font(name="Arial", size=10, bold=True, color=ROUGE)
    note(ws, 45, "Aucun chiffre de ce classeur ne représente une dépense constatée tant que les factures, la paie et le solde bancaire n'ont pas été rapprochés. "
                 "Trois hypothèses pèsent sur tout le reste : le taux de cotisation, le plafond et le taux de change. Elles fixent le coût d'une position, donc toute marge annoncée.", 4)
    ws.freeze_panes = "A5"
    return L


def onglet_grille(wb, m):
    ws = wb.create_sheet("Grille")
    logo(ws)
    titre(ws, "Grille tarifaire", "PRICING.md §1 fait foi. Tranches NON rétroactives depuis le 2026-09-18.")

    texte(ws.cell(row=4, column=1, value="Positions au tarif d'entrée"), taille=10)
    saisie(ws.cell(row=5, column=2, value=m["entree"]), NB)
    ws.cell(row=4, column=1).value = "Positions facturées au tarif d'entrée"
    ws.cell(row=5, column=1).value = "Tranche d'entrée"
    texte(ws.cell(row=5, column=1), taille=10)
    texte(ws.cell(row=5, column=3, value="positions"), taille=9, couleur="777777")

    entete(ws, 7, ["Métier", "Tarif d'entrée", "Tarif volume", "Commentaire"], [26, 15, 15, 76])
    for i, (cle, lab) in enumerate([("helpdesk", "helpdesk"), ("support", "support")]):
        r = 8 + i
        texte(ws.cell(row=r, column=1, value=lab), taille=10)
        saisie(ws.cell(row=r, column=2, value=m["offers"][cle]["mb"][0]), EUR)
        saisie(ws.cell(row=r, column=3, value=m["offers"][cle]["mb"][1]), EUR)
        texte(ws.cell(row=r, column=4, value="Positions 1 à 4 au tarif d'entrée, 5 et suivantes au tarif volume. La remise ne s'applique jamais aux positions déjà tarifées."),
              taille=9, couleur="777777", wrap=True)

    bloc(ws, 11, "PLANCHERS")
    texte(ws.cell(row=12, column=1, value="Plancher de négociation PRICING.md"), taille=10)
    saisie(ws.cell(row=12, column=2, value=m["plancher"]), EUR)
    texte(ws.cell(row=12, column=4, value="À RÉVISER. Garde-fou de négociation, pas seuil de rentabilité : comparer à la ligne 13, puis chiffrer le contrat avec la structure du lancement."),
          taille=9, couleur=ROUGE, wrap=True)
    texte(ws.cell(row=13, column=1, value="Prix minimal viable, quel que soit le volume"), taille=10)
    calcul(ws.cell(row=13, column=2, value="=(Hypothèses!B14+Hypothèses!B15/Hypothèses!B16)/(1-Hypothèses!B34-Hypothèses!B43)"), EUR2, gras=True,
           valeur=(m["agent"] + m["manager"] / m["a"]["managerCapacity"]) / (1 - m["reserveRate"]))
    texte(ws.cell(row=13, column=4, value="Coût d'un agent plus sa quote-part de manager, réserves sur CA comprises. En dessous, aucun volume ne rattrape le prix. Structure du lancement non comprise."),
          taille=9, couleur="777777", wrap=True)

    bloc(ws, 15, "CLIENT FINAL — AUCUNE GRILLE ARBITRÉE")
    for i, (cle, lab) in enumerate([("helpdesk", "helpdesk"), ("support", "support")]):
        r = 16 + i
        texte(ws.cell(row=r, column=1, value="Repère archivé " + lab), taille=10, couleur="777777")
        c = ws.cell(row=r, column=2, value=m["direct"][cle][0]); calcul(c, EUR); c.font = Font(name="Arial", size=10, color="777777", italic=True)
        c = ws.cell(row=r, column=3, value=m["direct"][cle][1]); calcul(c, EUR); c.font = Font(name="Arial", size=10, color="777777", italic=True)
    note(ws, 18, "Ces fourchettes viennent du registre interne archivé (évaluation du 2026-06-10). AUCUNE DÉCISION DATÉE NE LES FIXE, et les sources de cette étude sont "
                 "des sites de télésecrétariat médical et des BPO généralistes — aucune ne publie un prix de position dédiée en IT. Ne pas s'en servir comme tarif : "
                 "un entrant se chiffre au coût réel et à la marge visée. Détail : AUDIT-CHIFFRES-2026-09-18.md.", 4)
    ws.freeze_panes = "A8"


def onglet_contrats(wb, m, deals):
    ws = wb.create_sheet("Contrats")
    logo(ws)
    titre(ws, "Contrats", "Cellules bleues : à modifier. Une ligne par contrat, huit au maximum.")
    entete(ws, 5, ["Nom", "Mois de signature", "Métier", "Positions fermes", "Tarif", "Prix libre (€)",
                   "CA mensuel", "Onboarding", "Dépôt encaissé", "Mise en service", "Crédit de dépôt / facture", "Postes de travail"],
           [22, 12, 14, 12, 10, 12, 14, 12, 13, 13, 14, 13])

    for i in range(8):
        r = 6 + i
        d = deals[i] if i < len(deals) else None
        c = ws.cell(row=r, column=1, value=(d.get("label") or "") if d else "")
        c.font = Font(name="Arial", size=10, color=BLEU_SAISIE)
        saisie(ws.cell(row=r, column=2, value=d["signedMonth"] if d else None), NB)
        c = ws.cell(row=r, column=3, value=d["metier"] if d else None)
        c.font = Font(name="Arial", size=10, color=BLEU_SAISIE, bold=True)
        c.alignment = Alignment(horizontal="center")
        saisie(ws.cell(row=r, column=4, value=d["positions"] if d else None), NB)
        c = ws.cell(row=r, column=5, value=(d.get("tarif") or "mb") if d else None)
        c.font = Font(name="Arial", size=10, color=BLEU_SAISIE, bold=True)
        c.alignment = Alignment(horizontal="center")
        saisie(ws.cell(row=r, column=6, value=d.get("price") if d else None), EUR)

        calcul(ws.cell(row=r, column=7, value=(
            f'=IF($D{r}="",0,IF($E{r}="mb",'
            f'MIN($D{r},Grille!$B$5)*INDEX(Grille!$B$8:$B$9,MATCH($C{r},Grille!$A$8:$A$9,0))'
            f'+MAX(0,$D{r}-Grille!$B$5)*INDEX(Grille!$C$8:$C$9,MATCH($C{r},Grille!$A$8:$A$9,0)),'
            f'$D{r}*$F{r}))')), EUR, gras=True, valeur=(m["ca"][i] if d else 0))
        calcul(ws.cell(row=r, column=8, value=f'=IF($D{r}="",0,Hypothèses!$B$21+$D{r}*Hypothèses!$B$22)'), EUR,
               valeur=(m["onb"][i] if d else 0))
        calcul(ws.cell(row=r, column=9, value=f'=IF($D{r}="",0,$D{r}*Hypothèses!$B$27)'), EUR,
               valeur=(d["positions"] * m["a"]["depositPerPosition"] if d else 0))
        calcul(ws.cell(row=r, column=10, value=f'=IF($D{r}="",0,Hypothèses!$B$31+$D{r}*Hypothèses!$B$32)'), EUR,
               valeur=(m["a"]["activationFeePerDeal"] + d["positions"] * m["a"]["activationFeePerPosition"] if d else 0))
        calcul(ws.cell(row=r, column=11, value=f'=IF($D{r}="",0,$D{r}*Hypothèses!$B$28)'), EUR,
               valeur=(d["positions"] * m["a"]["invoiceCreditPerPosition"] if d else 0))
        calcul(ws.cell(row=r, column=12, value=f'=IF($D{r}="",0,$D{r}*Hypothèses!$B$41)'), EUR,
               valeur=(d["positions"] * m["a"]["equipmentPerPosition"] if d else 0))

    note(ws, 15, "Métier : helpdesk ou support, en minuscules. Tarif : « mb » applique la grille par tranches de l'onglet Grille ; toute autre valeur applique le prix libre de la colonne F. "
                 "En client final, le prix est OBLIGATOIREMENT saisi en colonne F : aucune grille directe n'existe. "
                 "Ligne d'exemple : Partenaire A · 4 · helpdesk · 2 · mb · (vide) → 4 000 € de CA mensuel.", 12)
    ws.freeze_panes = "A6"


def onglet_flux(wb, m, mois=24):
    ws = wb.create_sheet("Flux")
    logo(ws)
    titre(ws, "Flux mensuels", "Tout est formule. Modifiez Hypothèses ou Contrats : ce tableau suit.")
    entete(ws, 5, ["Mois", "Agents", "Mgr", "Dépôt", "Activation", "Facturé", "Encaissé",
                   "Personnel", "Onboarding et équipement", "Structure et réserves", "Résultat", "Cash fin de mois",
                   "Positions helpdesk", "Positions support"],
           [7, 8, 7, 11, 11, 12, 12, 12, 12, 11, 12, 15, 12, 12])

    D = "Contrats!$B$6:$B$13"
    P = "Contrats!$D$6:$D$13"
    CA = "Contrats!$G$6:$G$13"
    ONB = "Contrats!$H$6:$H$13"
    DEP = "Contrats!$I$6:$I$13"
    ACT = "Contrats!$J$6:$J$13"
    CRED = "Contrats!$K$6:$K$13"
    MET = "Contrats!$C$6:$C$13"
    EQ = "Contrats!$L$6:$L$13"

    deals = m["scenarios"]["reference"]
    hire = m["a"]["hireDelayMonths"]

    def positions(mois_no, metier):
        return sum(d["positions"] for d in deals
                   if d["metier"] == metier and mois_no >= d["signedMonth"] + hire)

    for i in range(mois):
        r = 6 + i
        rw = m["rows"][i]
        mno = i + 1
        calcul(ws.cell(row=r, column=1, value=mno), NB)
        ws.cell(row=r, column=1).font = Font(name="Arial", size=10, bold=True, color=ENCRE)

        calcul(ws.cell(row=r, column=13, value=f'=SUMPRODUCT(({MET}="helpdesk")*($A{r}>={D}+Hypothèses!$B$23)*{P})'), NB, valeur=positions(mno, "helpdesk"))
        calcul(ws.cell(row=r, column=14, value=f'=SUMPRODUCT(({MET}="support")*($A{r}>={D}+Hypothèses!$B$23)*{P})'), NB,
               valeur=positions(mno, "support"))
        calcul(ws.cell(row=r, column=2, value=f'=$M{r}+$N{r}'), NB, valeur=rw["agents"])
        calcul(ws.cell(row=r, column=3, value=(
            f'=IF($B{r}<=Hypothèses!$B$17,0,'
            f'IF($M{r}=0,0,CEILING($M{r}/Hypothèses!$B$16,1))+IF($N{r}=0,0,CEILING($N{r}/Hypothèses!$B$16,1)))')), NB,
               valeur=rw["managers"])

        calcul(ws.cell(row=r, column=4, value=f'=SUMPRODUCT(({D}=$A{r})*{DEP})'), EUR, valeur=rw["deposit"])
        calcul(ws.cell(row=r, column=5, value=f'=SUMPRODUCT(({D}=$A{r})*{ACT})'), EUR, valeur=rw["activation"])
        calcul(ws.cell(row=r, column=6, value=(
            f'=SUMPRODUCT(($A{r}>={D}+Hypothèses!$B$24)*{CA})'
            f'-SUMPRODUCT(($A{r}-{D}-Hypothèses!$B$24>=0)*($A{r}-{D}-Hypothèses!$B$24<Hypothèses!$B$29)*{CRED})')), EUR, valeur=rw["invoiced"])
        calcul(ws.cell(row=r, column=7, value=(
            f'=(SUMPRODUCT(($A{r}>={D}+Hypothèses!$B$24+Hypothèses!$B$26)*{CA})'
            f'-SUMPRODUCT(($A{r}-Hypothèses!$B$26-{D}-Hypothèses!$B$24>=0)'
            f'*($A{r}-Hypothèses!$B$26-{D}-Hypothèses!$B$24<Hypothèses!$B$29)*{CRED}))'
            f'*(1-Hypothèses!$B$33)')), EUR, valeur=rw["receipt"])
        calcul(ws.cell(row=r, column=8, value=f'=-($B{r}*Hypothèses!$B$14+$C{r}*Hypothèses!$B$15)'), EUR, valeur=-rw["payroll"])
        calcul(ws.cell(row=r, column=9, value=(
            f'=-SUMPRODUCT(({D}=$A{r})*{ONB})-SUMPRODUCT(({D}=$A{r})*{EQ})'
            f'-IF(AND($B{r}>0,COUNTIF($B$6:$B{r},">0")=1),Hypothèses!$B$42,0)')), EUR,
               valeur=-(rw["onboarding"] + rw["equipment"]))
        calcul(ws.cell(row=r, column=10, value=(
            f'=-(Hypothèses!$B$19+Hypothèses!$B$20+IF($B{r}>0,Hypothèses!$B$40,0)'
            f'+$G{r}*(Hypothèses!$B$34+Hypothèses!$B$43))')), EUR, valeur=-rw["overhead"])
        calcul(ws.cell(row=r, column=11, value=f'=$G{r}+$E{r}+$H{r}+$I{r}+$J{r}'), EUR, valeur=rw["result"])
        if i == 0:
            calcul(ws.cell(row=r, column=12, value=f'=Hypothèses!$B$36+$D{r}+$K{r}'), EUR, gras=True, valeur=rw["cash"])
        else:
            calcul(ws.cell(row=r, column=12, value=f'=L{r-1}+$D{r}+$K{r}'), EUR, gras=True, valeur=rw["cash"])

        if i % 2:
            for col in range(1, 15):
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor="FAFAF7")

    for col in (13, 14):
        ws.column_dimensions[get_column_letter(col)].hidden = True

    note(ws, 6 + mois + 1,
         "Personnel, onboarding, équipement, structure et réserves sont négatifs : ce sont des sorties. La structure ne se paie qu'avec du personnel en poste ; "
         "le groupe et l'UPS sortent une fois, au premier mois avec du personnel. Le « Résultat » compte les frais d'activation, qui sont acquis, "
         "mais pas le dépôt, qui n'est qu'une avance rendue par crédit de facture — le dépôt n'apparaît que dans le cash. "
         "Les colonnes M et N (positions par métier) sont masquées : elles servent au calcul du nombre de managers, qui n'est pas mutualisé entre métiers.", 12)
    ws.freeze_panes = "B6"
    return mois


def onglet_contribution(wb, m, mois):
    ws = wb.create_sheet("Contribution")
    logo(ws)
    titre(ws, "Contribution à régime plein", "Tous les contrats actifs, hors onboarding, équipement, dépôt et impayés.")
    entete(ws, 5, ["Contrat", "Positions", "CA mensuel", "Coût des agents", "Contribution avant supervision"],
           [24, 12, 14, 16, 22])

    lignes = m["contrib"]["lines"]
    agent = m["agent"]
    for i in range(8):
        r = 6 + i
        L = lignes[i] if i < len(lignes) else None
        calcul(ws.cell(row=r, column=1, value=f'=IF(Contrats!$D{r}="","",IF(Contrats!$A{r}="",Contrats!$C{r},Contrats!$A{r}))'),
               valeur=((f"Client {'AB'[i]}" if i < 2 else f"Contrat {i + 1}") if L else None))
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
        calcul(ws.cell(row=r, column=2, value=f'=IF(Contrats!$D{r}="","",Contrats!$D{r})'), NB,
               valeur=(L["deal"]["positions"] if L else None))
        calcul(ws.cell(row=r, column=3, value=f'=IF(Contrats!$D{r}="","",Contrats!$G{r})'), EUR,
               valeur=(L["revenue"] if L else None))
        calcul(ws.cell(row=r, column=4, value=f'=IF(Contrats!$D{r}="","",-Contrats!$D{r}*Hypothèses!$B$14)'), EUR,
               valeur=(-L["agentCost"] if L else None))
        calcul(ws.cell(row=r, column=5, value=f'=IF(Contrats!$D{r}="","",$C{r}+$D{r})'), EUR, gras=True,
               valeur=(L["revenue"] - L["agentCost"] if L else None))

    r = 15
    texte(ws.cell(row=r, column=1, value="Total des contrats"), taille=10, gras=True)
    totaux = {
        2: sum(x["deal"]["positions"] for x in lignes),
        3: sum(x["revenue"] for x in lignes),
        4: -sum(x["agentCost"] for x in lignes),
        5: sum(x["revenue"] - x["agentCost"] for x in lignes),
    }
    for col, lettre in [(2, "B"), (3, "C"), (4, "D"), (5, "E")]:
        calcul(ws.cell(row=r, column=col, value=f'=SUM({lettre}6:{lettre}13)'), NB if col == 2 else EUR, gras=True,
               valeur=totaux[col])
    ws.cell(row=r, column=1).border = Border(top=Side("thin", color=ENCRE))

    texte(ws.cell(row=16, column=1, value="Réserves SLA, banque et change"), taille=10)
    calcul(ws.cell(row=16, column=5, value='=-C15*(Hypothèses!$B$34+Hypothèses!$B$43)'), EUR,
           valeur=-totaux[3] * m["reserveRate"])

    fin = 5 + mois
    texte(ws.cell(row=17, column=1, value="Supervision à régime plein"), taille=10)
    mgr_fin = m["rows"][mois - 1]["managers"]
    calcul(ws.cell(row=17, column=5, value=f'=-Flux!C{fin}*Hypothèses!$B$15'), EUR, valeur=-mgr_fin * m["manager"])
    calcul(ws.cell(row=17, column=2, value=f'=Flux!C{fin}'), NB, valeur=mgr_fin)
    texte(ws.cell(row=18, column=1, value="Outils et charges supplémentaires"), taille=10)
    calcul(ws.cell(row=18, column=5, value='=-(Hypothèses!$B$19+Hypothèses!$B$20)'), EUR,
           valeur=-(m["a"]["toolsPerMonth"] + m["a"]["otherMonthlyCost"]))
    texte(ws.cell(row=19, column=1, value="Structure du lancement"), taille=10)
    calcul(ws.cell(row=19, column=5, value='=-IF(B15>0,Hypothèses!$B$40,0)'), EUR,
           valeur=-(m["a"]["structureMonthly"] if totaux[2] > 0 else 0))

    texte(ws.cell(row=20, column=1, value="RÉSULTAT MENSUEL À RÉGIME PLEIN"), taille=11, gras=True, couleur=ENCRE)
    c = ws.cell(row=20, column=5, value="=E15+E16+E17+E18+E19")
    c.font = Font(name="Arial", size=12, bold=True, color=ENCRE)
    c.number_format = EUR
    VALEURS.setdefault("Contribution", {})["E20"] = m["contrib"]["monthlyResult"]
    c.fill = PatternFill("solid", fgColor="EAF0EA")
    ws.cell(row=20, column=1).fill = PatternFill("solid", fgColor="EAF0EA")

    note(ws, 22, "La supervision est comptée globalement, au nombre de managers du dernier mois : elle n'est pas répartie par contrat, parce qu'un manager déclenché par "
                 "un contrat sert aussi au suivant dans le même métier. C'est le même « résultat mensuel à régime plein » que la page projection-finances-salverys.html "
                 "— CA moins réserves, agents, supervision, outils et structure.", 5)
    ws.freeze_panes = "A6"


def onglet_synthese(wb, m, mois):
    ws = wb.create_sheet("Synthèse", 0)
    logo(ws)
    titre(ws, "Projection de finances Salverys", "Prévision, pas situation bancaire. Chaque cellule calculée est une formule vivante.")
    fin = 5 + mois

    entete(ws, 5, ["Indicateur", "Valeur", "Lecture"], [38, 18, 78])
    lignes = [
        ("Creux de trésorerie", f'=MIN(Flux!L6:L{fin})', "Le plus bas du cash sur l'horizon. C'est le montant à financer si la trésorerie de départ est nulle.", EUR),
        ("Mois du creux", f'=INDEX(Flux!A6:A{fin},MATCH(MIN(Flux!L6:L{fin}),Flux!L6:L{fin},0))', "Quand il tombe.", NB),
        ("Besoin illustratif, creux + 30 %", f'=MAX(0,-MIN(Flux!L6:L{fin}))*1.3', "Coussin arbitraire de 30 %. Hors dettes et coûts non renseignés. PAS une recommandation de capital.", EUR),
        (f"Trésorerie à M{mois}", f'=Flux!L{fin}', "Fin de l'horizon, si aucun nouveau contrat n'est signé après ceux de l'onglet Contrats.", EUR),
        ("Résultat mensuel à régime plein", "=Contribution!E20", "Tous contrats actifs, structure et réserves comprises, hors onboarding, équipement, impayés et taxes.", EUR),
        ("Agents à régime plein", f'=Flux!B{fin}', "", NB),
        ("Managers à régime plein", f'=Flux!C{fin}', "Un par tranche de 8 positions, dans chaque métier. La supervision n'est pas mutualisée entre métiers.", NB),
        ("Coût d'une position", "=Hypothèses!B14", "Coût complet agent tout compris + VoIP (arbitrage A7, à valider par la paie).", EUR2),
        ("Prix minimal viable par position", "=Grille!B13", "En dessous, aucun volume ne rattrape le prix. Structure du lancement non comprise.", EUR2),
    ]
    cash = [r["cash"] for r in m["rows"]]
    creux = min(cash)
    vals = [creux, cash.index(creux) + 1, max(0, -creux) * 1.3, cash[-1],
            m["contrib"]["monthlyResult"], m["rows"][-1]["agents"], m["rows"][-1]["managers"],
            m["agent"], (m["agent"] + m["manager"] / m["a"]["managerCapacity"]) / (1 - m["reserveRate"])]
    for i, (lab, f, lecture, fmt) in enumerate(lignes):
        r = 6 + i
        texte(ws.cell(row=r, column=1, value=lab), taille=10, gras=(i == 0))
        c = ws.cell(row=r, column=2, value=f)
        c.font = Font(name="Arial", size=11, bold=True, color=ENCRE)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
        VALEURS.setdefault("Synthèse", {})[c.coordinate] = vals[i]
        texte(ws.cell(row=r, column=3, value=lecture), taille=9, couleur="777777", wrap=True)
        ws.row_dimensions[r].height = 24

    bloc(ws, 16, "CE QUE CE CLASSEUR NE CONTIENT PAS")
    note(ws, 17, "Impôts et TVA, assurance, amortissements comptables, recrutement au-delà du forfait d'onboarding, indemnités de rupture, dettes et créances existantes, dates réelles des factures annuelles. "
                 "Le modèle A mutualisé (1 120 € / 400 tickets) et le télésecrétariat médical facturé à l'appel ne sont pas projetés : leur staffing n'est pas mesuré.", 3)

    bloc(ws, 19, "COMMENT S'EN SERVIR")
    note(ws, 20, "Onglet Hypothèses : les cellules bleues pilotent tout le classeur. Onglet Contrats : une ligne par contrat, huit au maximum. "
                 "Les onglets Flux, Contribution et Synthèse sont entièrement calculés — n'y écrivez pas. "
                 "Ce classeur est généré par tools/export-xlsx.py depuis tools/finance-model.mjs : le régénérer écrase vos saisies, alors travaillez sur une copie.", 3)

    bloc(ws, 22, "OÙ EST LA VÉRITÉ")
    note(ws, 23, "PRICING.md fait foi sur les prix. AUDIT-CHIFFRES-2026-09-18.md dit, pour chaque chiffre du dépôt, s'il est vérifié, hypothétique, estimé ou faux aujourd'hui. "
                 "projection-finances-salverys.html est la version interactive de ce classeur ; grille-negociation-salverys.html dit à quel prix signer. "
                 "Les trois partagent le moteur tools/finance-model.mjs.", 3)
    ws.freeze_panes = "A6"


# -------------------------------------------------------------------------- main

def main():
    m = modele()
    deals = [dict(d, label=f"Client {'AB'[i]}") for i, d in enumerate(m["scenarios"]["reference"])]

    wb = Workbook()
    wb.remove(wb.active)
    onglet_hypotheses(wb, m)
    onglet_grille(wb, m)
    onglet_contrats(wb, m, deals)
    mois = onglet_flux(wb, m)
    onglet_contribution(wb, m, mois)
    onglet_synthese(wb, m, mois)

    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = ENCRE

    # Excel recalcule tout à l'ouverture : les valeurs injectées ne sont qu'un cache
    # pour les lecteurs qui ne recalculent pas.
    wb.calculation.fullCalcOnLoad = True
    wb.save(SORTIE)
    poses = injecter_valeurs(SORTIE, wb)
    taille = SORTIE.stat().st_size / 1024
    print(f"✓ {SORTIE.relative_to(ROOT)} — {len(wb.sheetnames)} onglets : {', '.join(wb.sheetnames)}")
    print(f"  {poses} valeurs mises en cache à côté des formules · {taille:.0f} Ko")


if __name__ == "__main__":
    main()
